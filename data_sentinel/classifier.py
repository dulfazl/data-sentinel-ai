"""Rule-based data classifier: file-name rules + content regexes validated with Luhn.

Levels: 0 Public, 1 Internal, 2 Confidential, 3 Restricted (highest match wins).
Only the *type* and *count* of each match is ever returned - never the matched value.
"""
import io
import os
import re
from dataclasses import dataclass, field

from . import config

LEVELS = config.LEVELS

TEXT_EXT = {".txt", ".csv", ".json", ".md", ".log", ".py", ".env", ".sql", ".xml",
            ".yaml", ".yml", ".ini", ".cfg", ".html"}
RICH_MAX_BYTES = 50 * 1024 * 1024          # .docx/.pdf larger than this are judged by name


def luhn_ok(num):
    digits = [int(c) for c in re.sub(r"\D", "", num)]
    if not 13 <= len(digits) <= 19:
        return False
    total = 0
    for i, d in enumerate(reversed(digits)):
        if i % 2:
            d = d * 2 - 9 if d * 2 > 9 else d * 2
        total += d
    return total % 10 == 0


# (label, compiled regex, level, min matches, optional validator)
CONTENT_RULES = [
    ("credit card number", re.compile(r"\b\d(?:[ -]?\d){12,18}\b"), 3, 1, luhn_ok),
    # Lookarounds stop it firing on 12 digits taken from a longer card number.
    ("Aadhaar-style ID", re.compile(r"(?<!\d)(?<!\d[\s-])\d{4}\s\d{4}\s\d{4}(?![\s-]?\d)"), 3, 1, None),
    ("PAN number", re.compile(r"\b[A-Z]{5}\d{4}[A-Z]\b"), 3, 1, None),
    ("US SSN", re.compile(r"\b\d{3}-\d{2}-\d{4}\b"), 3, 1, None),
    ("private key", re.compile(r"-----BEGIN (?:RSA |EC |DSA |OPENSSH |ENCRYPTED |PGP )?PRIVATE KEY(?: BLOCK)?-----"),
     3, 1, None),
    ("AWS access key", re.compile(r"\bAKIA[0-9A-Z]{16}\b"), 3, 1, None),
    ("hard-coded secret", re.compile(r"(?i)\b(?:password|passwd|secret|api[_-]?key)[\"']?\s*[:=]\s*[\"']?[^\s\"',;]{6,}"),
     3, 1, None),
    ("restricted marking", re.compile(r"(?i)\b(?:top secret|strictly confidential|restricted)\b"), 3, 1, None),
    ("confidential marking", re.compile(r"(?i)\bconfidential\b"), 2, 1, None),
    ("financial/HR terms", re.compile(r"(?i)\b(?:payroll|salary|salaries|customer list|balance sheet)\b"), 2, 1, None),
    ("bulk email addresses", re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+"), 2, 10, None),
    ("internal marking", re.compile(r"(?i)\binternal use only\b"), 1, 1, None),
]

FILENAME_RULES = [
    (re.compile(r"\.(pem|key|pfx|p12)$|(^|\.)env$|^\.env\.|id_rsa"), 3, "credential-type filename"),
    (re.compile(r"payroll|salary|passport|aadhaar|customer[_ -]?data"), 2, "sensitive filename keyword"),
    (re.compile(r"confidential|restricted|secret"), 2, "sensitivity keyword in filename"),
    (re.compile(r"\.(sql|bak|dump)$"), 2, "database/backup file type"),
]


@dataclass
class Classification:
    level: int = 0
    reasons: list = field(default_factory=list)
    matches: dict = field(default_factory=dict)      # label -> count (never values)
    inspected: bool = False                          # was the content read?
    note: str = ""                                   # why the content was not read

    @property
    def label(self):
        return LEVELS[self.level]


def classify_name(name, result=None):
    result = result or Classification()
    base = os.path.basename(name or "").lower()
    for rx, lvl, why in FILENAME_RULES:
        if rx.search(base):
            result.level = max(result.level, lvl)
            result.reasons.append(f"{why} -> {LEVELS[lvl]}")
    return result


def classify_text(text, result=None):
    result = result or Classification()
    result.inspected = True
    if not text:
        return result
    for label, rx, lvl, min_count, validator in CONTENT_RULES:
        hits = [m.group() for m in rx.finditer(text)]
        if validator:
            hits = [h for h in hits if validator(h)]
        if len(hits) >= min_count:
            result.level = max(result.level, lvl)
            result.matches[label] = len(hits)
            # never log the matched value itself, only its type and count
            result.reasons.append(f"content: {label} x{len(hits)} -> {LEVELS[lvl]}")
    return result


def classify(name, text=None, note=""):
    """Classify by file name, then by content when text is available."""
    result = classify_name(name)
    if text is None:
        result.note = note or "file type not readable"
        return result
    return classify_text(text, result)


# --------------------------------------------------------------------------- #
# Text extraction (in memory - no temp files)
# --------------------------------------------------------------------------- #
def _decode(data):
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="ignore")
    return data.decode("utf-8", errors="ignore")


def _docx_text(data):
    import docx  # python-docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(",".join(c.text for c in row.cells))
    return "\n".join(parts)


def _pdf_text(data, max_pages=50):
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages[:max_pages])


def can_read(name):
    ext = os.path.splitext(name or "")[1].lower()
    base = os.path.basename(name or "").lower()
    return ext in TEXT_EXT or base.startswith(".env") or ext in (".docx", ".pdf")


def extract_text(name, data):
    """Return (text or None, note). ``data`` is the file content as bytes."""
    ext = os.path.splitext(name or "")[1].lower()
    base = os.path.basename(name or "").lower()
    try:
        if ext in TEXT_EXT or base.startswith(".env"):
            return _decode(data[: config.TEXT_READ_LIMIT * 2])[: config.TEXT_READ_LIMIT], ""
        if ext == ".docx":
            try:
                return _docx_text(data)[: config.TEXT_READ_LIMIT], ""
            except ImportError:
                return None, "python-docx not installed"
        if ext == ".pdf":
            try:
                return _pdf_text(data)[: config.TEXT_READ_LIMIT], ""
            except ImportError:
                return None, "pdfplumber not installed"
    except Exception as exc:          # corrupt document: fall back to name-only
        return None, f"could not read content ({type(exc).__name__})"
    return None, f"file type '{ext or 'none'}' is judged by name only"


def read_limit(name):
    """How many bytes of a file on disk we need to read for classification."""
    ext = os.path.splitext(name or "")[1].lower()
    if ext in (".docx", ".pdf"):
        return RICH_MAX_BYTES
    return config.TEXT_READ_LIMIT * 2
