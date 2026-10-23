"""Build the IEEE-format report as PDF (two-column, via Chromium) and DOCX (editable, via python-docx).

Run:  .venv/bin/python report/make_figures.py && .venv/bin/python report/build_report.py
Out:  report/Data_Sentinel_AI_IEEE_Report.pdf, report/Data_Sentinel_AI_IEEE_Report.docx
"""
import html
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import content as C  # noqa: E402

FIG = os.path.join(HERE, "figures")
OUT = os.path.join(HERE, "Data_Sentinel_AI_IEEE_Report")
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
LETTERS = "ABCDEFGHIJKLMNOP"


# --------------------------------------------------------------------------- citations
def all_text():
    for block in C.BODY:
        for part in block[1:]:
            if isinstance(part, str):
                yield part
            elif isinstance(part, list):
                for item in part:
                    if isinstance(item, str):
                        yield item
                    elif isinstance(item, list):
                        yield from (x for x in item if isinstance(x, str))


ORDER = []
for text in all_text():
    for group in re.findall(r"\{ref:([^}]+)\}", text):
        for key in group.split(","):
            if key not in ORDER:
                ORDER.append(key)
missing = [k for k in C.REFS if k not in ORDER]
assert not missing, f"references never cited: {missing}"
NUM = {k: i + 1 for i, k in enumerate(ORDER)}


def cite(match):
    nums = sorted(NUM[k] for k in match.group(1).split(","))
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            out.append(f"[{nums[i]}]–[{nums[j]}]")
        else:
            out += [f"[{n}]" for n in nums[i:j + 1]]
        i = j + 1
    return ", ".join(out)


def resolve(text):
    return re.sub(r"\{ref:([^}]+)\}", cite, text)


def ref_text(r):
    a = r["authors"]
    if len(a) == 2 and a[1] == "et al.":
        names = f"{a[0]} <i>et al.</i>"
    elif len(a) == 1:
        names = a[0]
    elif len(a) == 2:
        names = f"{a[0]} and {a[1]}"
    else:
        names = ", ".join(a[:-1]) + ", and " + a[-1]
    vol = f"vol. {r['vol']}" + (f", no. {r['no']}" if r.get("no") else "")
    return (f"{names}, “{r['title']},” <i>{html.escape(r['journal'])}</i>, {vol}, pp. {r['pages']}, "
            f"{r['date']}, doi: {r['doi']}.")


# --------------------------------------------------------------------------- numbering pass
def numbered_blocks():
    h1 = h2 = h3 = fig = tab = 0
    out = []
    for b in C.BODY:
        kind = b[0]
        if kind == "h1":
            h1 += 1
            h2 = h3 = 0
            out.append(("h1", f"{ROMAN[h1 - 1]}.", b[1]))
        elif kind == "h2":
            h2 += 1
            h3 = 0
            out.append(("h2", f"{LETTERS[h2 - 1]}.", b[1]))
        elif kind == "h3":
            h3 += 1
            out.append(("h3", f"{h3})", b[1]))
        elif kind in ("fig", "code"):
            fig += 1
            out.append((kind, fig) + tuple(b[1:]))
        elif kind == "table":
            tab += 1
            out.append(("table", tab) + tuple(b[1:]))
        else:
            out.append(b)
    return out


BLOCKS = numbered_blocks()


# =========================================================================== PDF (HTML + Chromium)
CSS = """
@page { size: A4; margin: 19mm 14.3mm 20mm 14.3mm; }
* { box-sizing: border-box; }
body { font-family: "Times New Roman", Times, serif; font-size: 10pt; line-height: 1.17; color: #000;
       margin: 0; -webkit-print-color-adjust: exact; }
.title { font-size: 22pt; text-align: center; line-height: 1.12; margin: 0 0 12pt; font-weight: normal; }
.authors { display: flex; justify-content: center; gap: 30pt; text-align: center; margin-bottom: 7pt; }
.author .name { font-size: 11pt; } .author .roll { font-size: 9pt; font-style: italic; }
.meta { text-align: center; font-size: 9.5pt; margin: 0 0 16pt; line-height: 1.35; }
.cols { column-count: 2; column-gap: 4.22mm; text-align: justify; hyphens: auto; -webkit-hyphens: auto; orphans: 2; widows: 2; }
.abstract, .keywords { font-size: 9pt; font-weight: bold; margin: 0 0 7pt; text-indent: 1em; line-height: 1.16; }
p { margin: 0; text-indent: 1em; }
h1 { font-size: 10pt; font-weight: normal; font-variant: small-caps; text-align: center; margin: 11pt 0 5pt;
     letter-spacing: .03em; break-after: avoid; }
h2 { font-size: 10pt; font-weight: normal; font-style: italic; margin: 7pt 0 3pt; break-after: avoid; }
h3 { font-size: 10pt; font-weight: normal; font-style: italic; margin: 5pt 0 2pt 1em; break-after: avoid; }
ul, ol { margin: 2pt 0 2pt 0; padding-left: 1.6em; } li { margin: 0 0 1pt; }
figure { margin: 8pt 0 8pt; break-inside: avoid; text-align: center; }
figure img { max-width: 100%; }
figure.span { column-span: all; margin: 4pt 0 10pt; }
figcaption { font-size: 8pt; text-align: justify; margin-top: 4pt; line-height: 1.15; }
figure.span figcaption { text-align: center; }
.tbl { break-inside: avoid; margin: 8pt 0 8pt; }
.tbl.span { column-span: all; margin: 4pt 0 10pt; }
.tcap { text-align: center; font-size: 8pt; font-variant: small-caps; margin-bottom: 3pt; line-height: 1.2; }
table { border-collapse: collapse; width: 100%; font-size: 7.4pt; line-height: 1.15; text-align: left; }
th { border-top: 0.9pt solid #000; border-bottom: 0.5pt solid #000; font-weight: bold; padding: 2pt 3pt; vertical-align: bottom; }
td { padding: 1.6pt 3pt; vertical-align: top; }
tbody tr:last-child td { border-bottom: 0.9pt solid #000; }
.algo { border-top: 0.9pt solid #000; border-bottom: 0.9pt solid #000; margin: 8pt 0; break-inside: avoid; text-align: left; }
.algo .head { font-size: 8pt; font-weight: bold; border-bottom: 0.5pt solid #000; padding: 2pt 0; }
pre { font-family: "Courier New", Courier, monospace; font-size: 6.45pt; line-height: 1.22; margin: 3pt 0 3pt; white-space: pre-wrap; }
.code { border: 0.6pt solid #888; padding: 3pt 4pt; text-align: left; background: #fafafa; }
.eq { position: relative; text-align: center; margin: 5pt 0; padding: 0 22pt 0 0; font-size: 9.6pt; font-style: italic; text-indent: 0; }
.eq .num { position: absolute; right: 0; font-style: normal; }
.refs { font-size: 8pt; line-height: 1.17; }
.ref { display: flex; gap: 4pt; margin-bottom: 2.2pt; text-align: left; }
.ref .n { min-width: 16pt; flex: none; }
.ack { margin-bottom: 2pt; }
"""


def inline_html(text):
    return resolve(text)


def html_doc():
    h = [f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(C.TITLE)}</title>"
         f"<style>{CSS}</style></head><body>"]
    h.append(f"<div class='title'>{html.escape(C.TITLE)}</div>")
    h.append("<div class='authors'>" + "".join(
        f"<div class='author'><div class='name'>{html.escape(n)}</div><div class='roll'>{r}</div></div>"
        for n, r in C.AUTHORS) + "</div>")
    h.append(f"<div class='meta'>Guide: {html.escape(C.GUIDE)}<br>Course: {html.escape(C.COURSE)}</div>")
    h.append("<div class='cols'>")
    h.append(f"<p class='abstract'><i>Abstract</i>—{html.escape(C.ABSTRACT)}</p>")
    h.append(f"<p class='keywords'><i>Index Terms</i>—{html.escape(', '.join(C.KEYWORDS))}.</p>")
    for b in BLOCKS:
        kind = b[0]
        if kind == "h1":
            h.append(f"<h1>{b[1]} {html.escape(b[2])}</h1>")
        elif kind == "h2":
            h.append(f"<h2>{b[1]} {html.escape(b[2])}</h2>")
        elif kind == "h3":
            h.append(f"<h3>{b[1]} {html.escape(b[2])}</h3>")
        elif kind == "p":
            h.append(f"<p>{inline_html(b[1])}</p>")
        elif kind in ("bullets", "numbers"):
            tag = "ul" if kind == "bullets" else "ol"
            h.append(f"<{tag}>" + "".join(f"<li>{inline_html(i)}</li>" for i in b[1]) + f"</{tag}>")
        elif kind == "eq":
            h.append(f"<div class='eq'>{b[1]}<span class='num'>({b[2]})</span></div>")
        elif kind == "fig":
            n, name, cap, width, span = b[1], b[2], b[3], b[4], b[5]
            src = "file://" + os.path.join(FIG, name)
            h.append(f"<figure class='{'span' if span else ''}'><img src='{src}' style='width:{width}in'>"
                     f"<figcaption>Fig. {n}.&nbsp; {inline_html(cap)}</figcaption></figure>")
        elif kind == "code":
            n, cap, text = b[1], b[2], b[3]
            h.append(f"<figure><div class='code'><pre>{html.escape(text)}</pre></div>"
                     f"<figcaption>Fig. {n}.&nbsp; {inline_html(cap)}</figcaption></figure>")
        elif kind == "table":
            n, cap, head, rows, widths, span = b[1:7]
            total = sum(widths)
            cols = "".join(f"<col style='width:{100 * w / total:.1f}%'>" for w in widths)
            body = "".join("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in r) + "</tr>" for r in rows)
            h.append(f"<div class='tbl {'span' if span else ''}'><div class='tcap'>Table {ROMAN[n - 1]}<br>"
                     f"{html.escape(cap)}</div><table><colgroup>{cols}</colgroup><thead><tr>"
                     + "".join(f"<th>{html.escape(x)}</th>" for x in head) + f"</tr></thead><tbody>{body}</tbody></table></div>")
        elif kind == "pagebreak":
            h.append("<div style='break-before: page; height: 0'></div>")
        elif kind == "algo":
            h.append(f"<div class='algo'><div class='head'>{html.escape(b[1])}</div><pre>{html.escape(b[2])}</pre></div>")
    h.append("<h1>Acknowledgment</h1>")
    h.append(f"<p class='ack'>{html.escape(C.ACK)}</p>")
    h.append("<h1>References</h1><div class='refs'>")
    for key in ORDER:
        h.append(f"<div class='ref'><span class='n'>[{NUM[key]}]</span><span>{ref_text(C.REFS[key])}</span></div>")
    h.append("</div></div></body></html>")
    return "\n".join(h)


def build_pdf():
    from playwright.sync_api import sync_playwright
    page_html = OUT + ".html"
    with open(page_html, "w") as fh:
        fh.write(html_doc())
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="chromium", headless=True)
        page = browser.new_page()
        page.goto("file://" + page_html)
        page.wait_for_load_state("networkidle")
        page.pdf(path=OUT + ".pdf", prefer_css_page_size=True, print_background=True)
        browser.close()
    os.remove(page_html)
    return OUT + ".pdf"


# =========================================================================== DOCX (python-docx)
def build_docx():
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt

    doc = Document()
    COLW = Inches(3.45)

    def set_cols(section, n):
        sp = section._sectPr
        cols = sp.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sp.append(cols)
        cols.set(qn("w:num"), str(n))
        cols.set(qn("w:space"), "240")

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(14.3)
    sec.top_margin, sec.bottom_margin = Mm(19), Mm(20)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10)
    pf = normal.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    for sname, size, italic, caps, align in [("Heading 1", 10, False, True, WD_ALIGN_PARAGRAPH.CENTER),
                                             ("Heading 2", 10, True, False, WD_ALIGN_PARAGRAPH.LEFT),
                                             ("Heading 3", 10, True, False, WD_ALIGN_PARAGRAPH.LEFT)]:
        st = doc.styles[sname]
        st.font.name, st.font.size, st.font.bold, st.font.italic = "Times New Roman", Pt(size), False, italic
        st.font.small_caps = caps
        st.font.color.rgb = None
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(att), "Times New Roman")
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(att)) is not None:
                del rf.attrib[qn(att)]
        st.paragraph_format.alignment = align
        st.paragraph_format.space_before = Pt(8 if sname == "Heading 1" else 5)
        st.paragraph_format.space_after = Pt(4 if sname == "Heading 1" else 2)
        st.paragraph_format.keep_with_next = True
        if sname == "Heading 3":
            st.paragraph_format.left_indent = Inches(0.14)

    TAG = re.compile(r"(</?(?:b|i|sub|sup)>)")

    def add_inline(par, text, size=None, bold=False, italic=False, font=None):
        state = {"b": bold, "i": italic, "sub": False, "sup": False}
        for part in TAG.split(resolve(text)):
            if not part:
                continue
            m = re.fullmatch(r"<(/?)(b|i|sub|sup)>", part)
            if m:
                state[m.group(2)] = not m.group(1)
                continue
            run = par.add_run(html.unescape(part))
            run.bold, run.italic = state["b"], state["i"]
            run.font.subscript, run.font.superscript = state["sub"], state["sup"]
            if size:
                run.font.size = Pt(size)
            if font:
                run.font.name = font
        return par

    def para(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=None, **kw):
        p = doc.add_paragraph()
        p.alignment = align
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.14)
        if text:
            add_inline(p, text, size=size, **kw)
        return p

    def borders(table, top=True, bottom=True, box=False):
        tbl_pr = table._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            on = (edge == "top" and top) or (edge == "bottom" and bottom) or (box and edge in ("left", "right"))
            el.set(qn("w:val"), "single" if on else "nil")
            if on:
                el.set(qn("w:sz"), "8")
                el.set(qn("w:color"), "000000")
            b.append(el)
        tbl_pr.append(b)

    def cell_bottom(cell):
        tc_pr = cell._tc.get_or_add_tcPr()
        tb = OxmlElement("w:tcBorders")
        el = OxmlElement("w:bottom")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        tb.append(el)
        tc_pr.append(tb)

    def cell_par(cell, text, size, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font=None):
        p = cell.paragraphs[0] if not cell.paragraphs[0].text and len(cell.paragraphs) == 1 and \
            not cell.paragraphs[0].runs else cell.add_paragraph()
        p.alignment = align
        p.paragraph_format.first_line_indent = Inches(0)
        add_inline(p, text, size=size, bold=bold, font=font)
        return p

    def caption(text, center=False):
        p = para(text, align=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY, indent=False,
                 size=8)
        p.paragraph_format.space_after = Pt(8)
        return p

    def new_section(cols):
        s = doc.add_section(WD_SECTION.CONTINUOUS)
        set_cols(s, cols)
        return s

    # ---------------------------------------------------------------- title block (1 column)
    t = para(C.TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=22)
    t.paragraph_format.space_after = Pt(12)
    ap = para("", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    for k, (name, roll) in enumerate(C.AUTHORS):
        r = ap.add_run(("   \u00b7   " if k else "") + name)
        r.font.size = Pt(11)
        r2 = ap.add_run(f" ({roll})")
        r2.font.size, r2.italic = Pt(9.5), True
    m = para(f"Guide: {C.GUIDE}", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=9.5)
    m.paragraph_format.space_before = Pt(6)
    m2 = para(f"Course: {C.COURSE}", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=9.5)
    m2.paragraph_format.space_after = Pt(12)
    new_section(2)

    p = para(indent=True)
    add_inline(p, "<i>Abstract</i>—" + html.escape(C.ABSTRACT), size=9, bold=True)
    p.paragraph_format.space_after = Pt(6)
    p = para(indent=True)
    add_inline(p, "<i>Index Terms</i>—" + ", ".join(C.KEYWORDS) + ".", size=9, bold=True)
    p.paragraph_format.space_after = Pt(4)

    for b in BLOCKS:
        kind = b[0]
        if kind in ("h1", "h2", "h3"):
            doc.add_paragraph(f"{b[1]} {b[2]}", style={"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}[kind])
        elif kind == "p":
            para(b[1])
        elif kind in ("bullets", "numbers"):
            for item in b[1]:
                lp = doc.add_paragraph(style="List Bullet" if kind == "bullets" else "List Number")
                lp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_inline(lp, item)
        elif kind == "eq":
            ep = doc.add_paragraph()
            ep.paragraph_format.tab_stops.add_tab_stop(Inches(1.72), WD_TAB_ALIGNMENT.CENTER)
            ep.paragraph_format.tab_stops.add_tab_stop(Inches(3.43), WD_TAB_ALIGNMENT.RIGHT)
            ep.paragraph_format.space_before = ep.paragraph_format.space_after = Pt(4)
            ep.add_run("\t")
            add_inline(ep, b[1], italic=True)
            ep.add_run(f"\t({b[2]})")
        elif kind == "fig":
            n, name, cap, width, span = b[1], b[2], b[3], b[4], b[5]
            if span:
                new_section(1)
            fp = doc.add_paragraph()
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fp.paragraph_format.space_before = Pt(6)
            fp.paragraph_format.keep_with_next = True
            fp.add_run().add_picture(os.path.join(FIG, name), width=Inches(min(width, 7.0 if span else 3.45)))
            caption(f"Fig. {n}. {cap}", center=span)
            if span:
                new_section(2)
        elif kind == "code":
            n, cap, text = b[1], b[2], b[3]
            tb = doc.add_table(rows=1, cols=1)
            borders(tb, box=True)
            cell = tb.cell(0, 0)
            cell.width = COLW
            for line in text.split("\n"):
                cp = cell_par(cell, html.escape(line) or " ", 6.3, font="Courier New")
            caption(f"Fig. {n}. {cap}")
        elif kind == "table":
            n, cap, head, rows, widths, span = b[1:7]
            if span:
                new_section(1)
            cp = para(f"Table {ROMAN[n - 1]}", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=8)
            cp.runs[0].font.small_caps = True
            cp.paragraph_format.space_before = Pt(6)
            cp.paragraph_format.keep_with_next = True
            cp2 = para(html.escape(cap), align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=8)
            cp2.runs[0].font.small_caps = True
            cp2.paragraph_format.keep_with_next = True
            cp2.paragraph_format.space_after = Pt(3)
            tbl = doc.add_table(rows=1 + len(rows), cols=len(head))
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            borders(tbl)
            full = 7.0 if span else 3.45
            scale = full / sum(widths)
            for j, h in enumerate(head):
                c = tbl.cell(0, j)
                c.width = Inches(widths[j] * scale)
                cell_par(c, html.escape(h), 7.5, bold=True)
                cell_bottom(c)
            for i, r in enumerate(rows, 1):
                for j, val in enumerate(r):
                    c = tbl.cell(i, j)
                    c.width = Inches(widths[j] * scale)
                    cell_par(c, val, 7.5)
            sp = para("", indent=False)
            sp.paragraph_format.space_after = Pt(4)
            if span:
                new_section(2)
        elif kind == "algo":
            tb = doc.add_table(rows=2, cols=1)
            borders(tb)
            head_cell, body_cell = tb.cell(0, 0), tb.cell(1, 0)
            head_cell.width = body_cell.width = COLW
            cell_par(head_cell, html.escape(b[1]), 8, bold=True)
            cell_bottom(head_cell)
            for line in b[2].split("\n"):
                cell_par(body_cell, html.escape(line) or " ", 6.3, font="Courier New")
            sp = para("", indent=False)
            sp.paragraph_format.space_after = Pt(4)

    doc.add_paragraph("Acknowledgment", style="Heading 1")
    para(html.escape(C.ACK))
    doc.add_paragraph("References", style="Heading 1")
    for key in ORDER:
        rp = doc.add_paragraph()
        rp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rp.paragraph_format.left_indent = Inches(0.3)
        rp.paragraph_format.first_line_indent = Inches(-0.3)
        rp.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
        rp.paragraph_format.space_after = Pt(2)
        add_inline(rp, f"[{NUM[key]}]\t" + ref_text(C.REFS[key]), size=8)
    new_section(2)                                  # balances the last page's columns
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = Mm(14.3)
        s.top_margin, s.bottom_margin = Mm(19), Mm(20)
    doc.core_properties.title = C.TITLE
    doc.core_properties.author = ", ".join(n for n, _ in C.AUTHORS)
    doc.save(OUT + ".docx")
    return OUT + ".docx"


if __name__ == "__main__":
    for b in C.BODY:
        if b[0] == "algo":
            for line in b[2].split("\n"):
                if len(line) > 64:
                    print(f"  warning: long pseudocode line ({len(line)}): {line}")
    print("references in citation order:", ", ".join(f"[{NUM[k]}] {k}" for k in ORDER))
    print(build_pdf())
    print(build_docx())
